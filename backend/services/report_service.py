import io
import os
from datetime import datetime, date
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, portrait
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.platypus import Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
def get_faculty_initials(name):
    if not name:
        return 'TBD'
    mapping = {
        'Dr. Arjun B C': 'ABC',
        'Dr. Swathi H Y': 'HYS',
        'Mrs. Sushma M V': 'SMV',
        'Mrs. Ankitha S': 'ANS',
        'Mrs. Megha H C': 'MHC',
        'Mrs. Maseeha Banu': 'MB',
        'Sheethal D N': 'SDN'
    }
    for k, v in mapping.items():
        if k.lower() in name.lower() or name.lower() in k.lower():
            return v
    clean_parts = [p for p in name.split() if p.upper() not in ('DR.', 'DR', 'MRS.', 'MRS', 'MR.', 'MR', 'PROF.', 'PROF', 'MS.', 'MS')]
    initials = ''.join([p[0].upper() for p in clean_parts if p])
    return initials or 'FAC'


def format_time_slot(st, et):
    if not st or not et:
        return f"{st} to {et}"
    try:
        sh, sm = map(int, st.split(':'))
        eh, em = map(int, et.split(':'))
        s_ampm = "AM" if sh < 12 else "PM"
        e_ampm = "AM" if eh < 12 else "PM"
        sh_12 = sh if sh <= 12 else sh - 12
        if sh_12 == 0: sh_12 = 12
        eh_12 = eh if eh <= 12 else eh - 12
        if eh_12 == 0: eh_12 = 12
        s_str = f"{sh_12}.{sm:02d} {s_ampm}" if sm != 0 else f"{sh_12}.00 {s_ampm}"
        e_str = f"{eh_12}.{em:02d} {e_ampm}" if em != 0 else f"{eh_12}.00 {e_ampm}"
        return f"{s_str} to {e_str}"
    except Exception:
        return f"{st} to {et}"


BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
CREST_PATH = os.path.join(BASE_DIR, 'mce_crest.png')
AIML_PATH = os.path.join(BASE_DIR, 'aiml_logo.png')


class ReportService:

    @staticmethod
    def _build_official_header_table(cie, doc_type_title, academic_term_str=""):
        """
        Creates the authentic institutional header table matching File 2:
        Left: MCE Crest | Center: College, Dept, CIE, and Document Titles | Right: AIML Logo
        """
        styles = getSampleStyleSheet()
        t_college = ParagraphStyle('HCol', parent=styles['Normal'], fontName='Times-Bold', fontSize=13, leading=16, alignment=1)
        t_dept = ParagraphStyle('HDept', parent=styles['Normal'], fontName='Times-Bold', fontSize=12.5, leading=15.5, alignment=1)
        t_aiml = ParagraphStyle('HAiml', parent=styles['Normal'], fontName='Times-Bold', fontSize=12.5, leading=15.5, alignment=1)
        t_cie = ParagraphStyle('HCIE', parent=styles['Normal'], fontName='Times-Bold', fontSize=13, leading=16, alignment=1)
        t_title = ParagraphStyle('HTitle', parent=styles['Normal'], fontName='Times-Bold', fontSize=13, leading=16, alignment=1)

        center_content = [
            Paragraph("<b>Malnad College of Engineering, Hassan</b>", t_college),
            Paragraph("<b>Department of Computer Science & Engineering</b>", t_dept),
            Paragraph("<b>(Artificial Intelligence & Machine Learning)</b>", t_aiml),
            Spacer(1, 2),
        ]

        if academic_term_str:
            # Page 1 or Page 2
            center_content.append(Paragraph(f"<b>{cie.name}</b>", t_cie))
            center_content.append(Paragraph(f"<b>{doc_type_title} {academic_term_str}</b>", t_title))
        else:
            # Page 3: Faculty-wise Allotment
            center_content.append(Paragraph(f"<b>{doc_type_title}</b>", t_title))

        left_img = Paragraph("", styles['Normal'])
        if os.path.exists(CREST_PATH):
            try:
                left_img = RLImage(CREST_PATH, width=50, height=46)
            except Exception:
                pass

        right_img = Paragraph("", styles['Normal'])
        if os.path.exists(AIML_PATH):
            try:
                right_img = RLImage(AIML_PATH, width=50, height=46)
            except Exception:
                pass

        header_table = Table([[left_img, center_content, right_img]], colWidths=[52, 418, 52])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ALIGN', (0, 0), (0, 0), 'LEFT'),
            ('ALIGN', (1, 0), (1, 0), 'CENTER'),
            ('ALIGN', (2, 0), (2, 0), 'RIGHT'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        return header_table

    @staticmethod
    def _build_three_signatories_table():
        """
        Creates the authentic 3-signatory block matching File 2:
        Left: (Mrs. Ankitha S) CIE Coordinator
        Center: (Dr. Swathi H Y) Chairman, BOE
        Right: (Dr. Arjun B C) HOD
        """
        styles = getSampleStyleSheet()
        sig_name_style = ParagraphStyle('SigName', parent=styles['Normal'], fontName='Times-Bold', fontSize=11, leading=13, alignment=1)
        sig_role_style = ParagraphStyle('SigRole', parent=styles['Normal'], fontName='Times-Roman', fontSize=10.5, leading=12.5, alignment=1)
        
        sig_data = [
            [
                Paragraph("<b>(Mrs. Ankitha S)</b>", sig_name_style),
                Paragraph("<b>(Dr. Swathi H Y)</b>", sig_name_style),
                Paragraph("<b>(Dr. Arjun B C)</b>", sig_name_style)
            ],
            [
                Paragraph("CIE Coordinator", sig_role_style),
                Paragraph("Chairman, BOE", sig_role_style),
                Paragraph("HOD", sig_role_style)
            ]
        ]
        sig_table = Table(sig_data, colWidths=[174, 174, 174])
        sig_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('TOPPADDING', (0, 0), (-1, -1), 1),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ]))
        return sig_table

    @staticmethod
    def generate_cie_pdf(cie, duties, selected_semester=None, semester_ids=None):
        """
        Generates the exact authentic 3-Page institutional CIE Examination Duty Allotment PDF matching File 2:
        - Page 1: Invigilation Duty Allotment (AY ...) Even/Odd (Grouped by Semester)
        - Page 2: Squad Duty Allotment (AY ...) Even/Odd (Grouped by Semester, rooms separated by gridlines)
        - Page 3: CIE Duty Allotment (Faculty-wise Roster with Sl No, Faculty Name, Date, Time, Duty, Signature)
        Each page features the official MCE Crest, AIML emblem, and 3 institutional signatories at the bottom.
        Strictly excludes any faculty not present in the active Faculty Directory.
        """
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=(595.28, 841.89), # A4 dimensions matching authentic institutional PDF
            leftMargin=36,
            rightMargin=36,
            topMargin=22,
            bottomMargin=22
        )

        elements = []
        styles = getSampleStyleSheet()
        sem_title_style = ParagraphStyle('SemT', parent=styles['Normal'], fontName='Times-Bold', fontSize=12, leading=15, alignment=1)
        cell_bold = ParagraphStyle('CB', parent=styles['Normal'], fontName='Times-Bold', fontSize=9.5, leading=11.5, alignment=1)
        cell_norm = ParagraphStyle('CN', parent=styles['Normal'], fontName='Times-Roman', fontSize=9.5, leading=11.5, alignment=1)

        # Determine academic term
        is_even = "even" in (cie.name or '').lower() or (cie.start_date and cie.start_date.month in (2, 3, 4, 5, 6))
        term_str = f"(AY {cie.academic_year or '2026-27'}) {'Even' if is_even else 'Odd'}"

        from models import ExamSession, Semester, Faculty
        # Strictly query active faculty directory IDs
        active_facs = Faculty.query.filter_by(is_active=True).all()
        active_dir_ids = {f.id for f in active_facs}

        # Fetch sessions for this CIE
        session_query = ExamSession.query.filter_by(cie_id=cie.id)
        if selected_semester:
            session_query = session_query.filter_by(semester_id=selected_semester.id)
        elif semester_ids:
            session_query = session_query.filter(ExamSession.semester_id.in_(semester_ids))
        else:
            # Exclude 3rd semester from combined view so that 5th & 7th remain combined as in the official circular
            session_query = session_query.join(Semester).filter(Semester.sem_number != 3)
        sessions = session_query.order_by(ExamSession.exam_date, ExamSession.start_time).all()

        # Group sessions by semester
        semesters_map = defaultdict(list)
        for s in sessions:
            sem_name = s.semester.name if s.semester else 'Academic Class'
            sem_number = s.semester.sem_number if s.semester else 0
            semesters_map[(sem_number, sem_name)].append(s)

        sorted_semesters = sorted(semesters_map.items(), key=lambda x: x[0][0])

        total_p1_rows = sum(len(s_list) * 2 for _, s_list in sorted_semesters)
        p1_pad = 4.5 if total_p1_rows <= 14 else 2.5

        # =========================================================================
        # PAGE 1: INVIGILATION DUTY ALLOTMENT
        # =========================================================================
        header_p1 = ReportService._build_official_header_table(cie, "Invigilation Duty Allotment", term_str)
        elements.append(header_p1)
        elements.append(Spacer(1, 4))

        p1_table_heights = []
        for (sem_num, sem_name), s_list in sorted_semesters:
            elements.append(Paragraph(f"<b>{sem_name}</b>", sem_title_style))
            elements.append(Spacer(1, 2))

            # Group sessions by date
            date_map = defaultdict(list)
            for s in s_list:
                date_map[s.exam_date].append(s)

            table_rows = [
                [
                    Paragraph("<b>Date</b>", cell_bold),
                    Paragraph("<b>Time</b>", cell_bold),
                    Paragraph("<b>Course Code</b>", cell_bold),
                    Paragraph("<b>Room No.</b>", cell_bold),
                    Paragraph("<b>Faculty</b>", cell_bold)
                ]
            ]
            table_styles = [
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), p1_pad),
                ('BOTTOMPADDING', (0, 0), (-1, -1), p1_pad),
                ('LEFTPADDING', (0, 0), (-1, -1), 2),
                ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ]

            current_row = 1
            for exam_date, d_sessions in sorted(date_map.items(), key=lambda x: x[0]):
                date_start_row = current_row
                date_p = Paragraph(f"<b>{exam_date.strftime('%d.%m.%Y')}</b><br/><b>{exam_date.strftime('%A')}</b>", cell_bold)

                for sess in d_sessions:
                    time_p = Paragraph(format_time_slot(sess.start_time, sess.end_time), cell_bold)
                    code_p = Paragraph(sess.subject.code if sess.subject else '', cell_norm)
                    session_start_row = current_row

                    # Strictly include duties of active directory faculties only
                    invig_duties = [d for d in sess.duties if (d.duty_type or '').lower() != 'squad duty' and d.faculty_id in active_dir_ids and d.faculty and d.faculty.is_active]
                    rooms_to_render = []

                    if sess.exam_rooms:
                        for er in sess.exam_rooms:
                            matching_d = next((d for d in invig_duties if d.exam_room_id == er.id), None)
                            if not matching_d:
                                matching_d = next((d for d in invig_duties if not d.exam_room_id and d.faculty_id), None)
                            fac_init = get_faculty_initials(matching_d.faculty.name if matching_d and matching_d.faculty else None)
                            rooms_to_render.append((er.room_number, fac_init))
                    elif invig_duties:
                        for d in invig_duties:
                            r_no = d.exam_room.room_number if d.exam_room else (sess.room_number or 'AI301')
                            fac_init = get_faculty_initials(d.faculty.name if d.faculty else None)
                            rooms_to_render.append((r_no, fac_init))
                    else:
                        rooms_to_render.append(('AI301', ''))
                        rooms_to_render.append(('AI302', ''))

                    for r_idx, (r_no, fac_init) in enumerate(rooms_to_render):
                        is_date_first = (current_row == date_start_row)
                        is_sess_first = (current_row == session_start_row)
                        table_rows.append([
                            date_p if is_date_first else '',
                            time_p if is_sess_first else '',
                            code_p if is_sess_first else '',
                            Paragraph(r_no, cell_norm),
                            Paragraph(fac_init, cell_norm)
                        ])
                        current_row += 1

                    session_end_row = current_row - 1
                    if session_end_row > session_start_row:
                        table_styles.append(('SPAN', (1, session_start_row), (1, session_end_row)))
                        table_styles.append(('SPAN', (2, session_start_row), (2, session_end_row)))

                date_end_row = current_row - 1
                if date_end_row > date_start_row:
                    table_styles.append(('SPAN', (0, date_start_row), (0, date_end_row)))

            t_invig = Table(table_rows, colWidths=[70, 130, 100, 111, 111])
            t_invig.setStyle(TableStyle(table_styles))
            elements.append(t_invig)
            elements.append(Spacer(1, 4))
            _, t_h = t_invig.wrap(522, 800)
            p1_table_heights.append(t_h)

        # Dynamic spacer calculation so signatories sit gracefully near the bottom (matching File 2)
        _, h_head1 = header_p1.wrap(522, 800)
        sig_p1 = ReportService._build_three_signatories_table()
        _, h_sig1 = sig_p1.wrap(522, 800)
        p1_used = h_head1 + 4 + sum(p1_table_heights) + len(sorted_semesters) * (15 + 2 + 4)
        p1_push = max(20, min(360, 740 - p1_used - h_sig1))
        elements.append(Spacer(1, p1_push))
        elements.append(sig_p1)

        # =========================================================================
        # PAGE 2: SQUAD DUTY ALLOTMENT
        # =========================================================================
        elements.append(PageBreak())
        header_p2 = ReportService._build_official_header_table(cie, "Squad Duty Allotment", term_str)
        elements.append(header_p2)
        elements.append(Spacer(1, 4))

        total_p2_rows = sum(len(s_list) * 2 for _, s_list in sorted_semesters)
        p2_pad = 4.5 if total_p2_rows <= 14 else 2.5
        p2_table_heights = []

        for (sem_num, sem_name), s_list in sorted_semesters:
            elements.append(Paragraph(f"<b>{sem_name}</b>", sem_title_style))
            elements.append(Spacer(1, 2))

            date_map = defaultdict(list)
            for s in s_list:
                date_map[s.exam_date].append(s)

            table_rows = [
                [
                    Paragraph("<b>Date</b>", cell_bold),
                    Paragraph("<b>Time</b>", cell_bold),
                    Paragraph("<b>Course Code</b>", cell_bold),
                    Paragraph("<b>Room No.</b>", cell_bold),
                    Paragraph("<b>Faculty</b>", cell_bold)
                ]
            ]
            table_styles = [
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), p2_pad),
                ('BOTTOMPADDING', (0, 0), (-1, -1), p2_pad),
                ('LEFTPADDING', (0, 0), (-1, -1), 2),
                ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ]

            current_row = 1
            for exam_date, d_sessions in sorted(date_map.items(), key=lambda x: x[0]):
                date_start_row = current_row
                date_p = Paragraph(f"<b>{exam_date.strftime('%d.%m.%Y')}</b><br/><b>{exam_date.strftime('%A')}</b>", cell_bold)

                for sess in d_sessions:
                    session_start_row = current_row
                    time_p = Paragraph(format_time_slot(sess.start_time, sess.end_time), cell_bold)
                    code_p = Paragraph(sess.subject.code if sess.subject else '', cell_norm)

                    # Squad duties filtered by active directory faculty only
                    squad_duties = [d for d in sess.duties if (d.duty_type or '').lower() == 'squad duty' and d.faculty_id in active_dir_ids and d.faculty and d.faculty.is_active]
                    squad_initials = [get_faculty_initials(sd.faculty.name) for sd in squad_duties]
                    squad_str = ', '.join(squad_initials)
                    fac_p = Paragraph(squad_str, cell_norm) if squad_str else Paragraph("", cell_norm)

                    # Render each room on its own row with grid lines (matching Page 1 and File 2)
                    rooms_list = [r.room_number for r in sess.exam_rooms] or [sess.room_number or 'AI301']
                    if len(rooms_list) == 1 and not sess.exam_rooms:
                        rooms_list = ['AI301', 'AI302']

                    for r_idx, r_no in enumerate(rooms_list):
                        is_date_first = (current_row == date_start_row)
                        is_sess_first = (current_row == session_start_row)

                        table_rows.append([
                            date_p if is_date_first else '',
                            time_p if is_sess_first else '',
                            code_p if is_sess_first else '',
                            Paragraph(r_no, cell_norm),
                            fac_p if is_sess_first else ''
                        ])
                        current_row += 1

                    session_end_row = current_row - 1
                    if session_end_row > session_start_row:
                        table_styles.append(('SPAN', (1, session_start_row), (1, session_end_row)))
                        table_styles.append(('SPAN', (2, session_start_row), (2, session_end_row)))
                        table_styles.append(('SPAN', (4, session_start_row), (4, session_end_row)))

                date_end_row = current_row - 1
                if date_end_row > date_start_row:
                    table_styles.append(('SPAN', (0, date_start_row), (0, date_end_row)))

            t_squad = Table(table_rows, colWidths=[70, 130, 100, 111, 111])
            t_squad.setStyle(TableStyle(table_styles))
            elements.append(t_squad)
            elements.append(Spacer(1, 4))
            _, t_h2 = t_squad.wrap(522, 800)
            p2_table_heights.append(t_h2)

        # Dynamic spacer calculation so signatories sit gracefully near the bottom on Page 2
        _, h_head2 = header_p2.wrap(522, 800)
        sig_p2 = ReportService._build_three_signatories_table()
        _, h_sig2 = sig_p2.wrap(522, 800)
        p2_used = h_head2 + 4 + sum(p2_table_heights) + len(sorted_semesters) * (15 + 2 + 4)
        p2_push = max(20, min(360, 740 - p2_used - h_sig2))
        elements.append(Spacer(1, p2_push))
        elements.append(sig_p2)

        # =========================================================================
        # PAGE 3: FACULTY-WISE DUTY ALLOTMENT & SIGN-OFF
        # (Grouped by Faculty, with dates consolidated and duty/signature spanned)
        # =========================================================================
        elements.append(PageBreak())
        header_p3 = ReportService._build_official_header_table(cie, f"{cie.name} Duty Allotment", "")
        elements.append(header_p3)
        elements.append(Spacer(1, 4))

        # Query all assigned duties
        from models import Duty, Semester
        duty_query = Duty.query.join(ExamSession).join(Semester).filter(ExamSession.cie_id == cie.id)
        if selected_semester:
            duty_query = duty_query.filter(ExamSession.semester_id == selected_semester.id)
        elif semester_ids:
            duty_query = duty_query.filter(ExamSession.semester_id.in_(semester_ids))
        else:
            duty_query = duty_query.filter(Semester.sem_number != 3)
        all_duties = duty_query.order_by(ExamSession.exam_date, ExamSession.start_time).all()

        faculty_duties_map = defaultdict(list)
        for d in all_duties:
            if d.faculty_id in active_dir_ids and d.faculty and d.faculty.is_active:
                faculty_duties_map[d.faculty.name].append(d)

        def fac_sort_key(name):
            if 'Arjun' in name: return 0
            if 'Swathi' in name: return 1
            if 'Sushma' in name: return 2
            if 'Ankitha' in name: return 3
            if 'Megha' in name: return 4
            if 'Maseeha' in name: return 5
            if 'Sheethal' in name: return 6
            return 9

        sorted_fac_entries = sorted(faculty_duties_map.items(), key=lambda x: (fac_sort_key(x[0]), x[0]))

        name_cell_style = ParagraphStyle('NameCell', parent=styles['Normal'], fontName='Times-Bold', fontSize=9.5, leading=12, alignment=0)
        p3_pad = 4.0 if len(sorted_fac_entries) <= 7 else 2.5

        f_table_rows = [
            [
                Paragraph("<b>Sl No.</b>", cell_bold),
                Paragraph("<b>Name of the Faculty</b>", cell_bold),
                Paragraph("<b>Date</b>", cell_bold),
                Paragraph("<b>Time</b>", cell_bold),
                Paragraph("<b>Duty</b>", cell_bold),
                Paragraph("<b>Signature</b>", cell_bold)
            ]
        ]
        f_table_styles = [
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), p3_pad),
            ('BOTTOMPADDING', (0, 0), (-1, -1), p3_pad),
            ('LEFTPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ('LEFTPADDING', (1, 0), (1, -1), 6), # Clean indent for faculty names
        ]

        cur_f_row = 1
        for idx, (fac_name, f_duties) in enumerate(sorted_fac_entries, 1):
            f_start = cur_f_row
            sl_p = Paragraph(f"<b>{idx}.</b>", cell_bold)
            name_p = Paragraph(f"<b>{fac_name}</b>", name_cell_style)

            # Separate Invigilation and Squad duties
            invig_duties = [d for d in f_duties if (d.duty_type or '').lower() != 'squad duty']
            squad_duties = [d for d in f_duties if (d.duty_type or '').lower() == 'squad duty']

            invig_by_date = defaultdict(list)
            for d in invig_duties:
                sess = d.exam_session
                invig_by_date[sess.exam_date].append((sess.start_time, sess.end_time))

            squad_by_date = defaultdict(list)
            for d in squad_duties:
                sess = d.exam_session
                squad_by_date[sess.exam_date].append((sess.start_time, sess.end_time))

            duty_blocks = []
            # 1. Invigilation block (all invigilations listed together)
            if invig_by_date:
                invig_entries = []
                for dt in sorted(invig_by_date.keys()):
                    sorted_slots = sorted(invig_by_date[dt], key=lambda x: x[0])
                    time_str = '<br/>'.join([format_time_slot(st, et) for st, et in sorted_slots])
                    invig_entries.append((dt.strftime('%d.%m.%Y'), time_str))
                duty_blocks.append(('Invigilation', invig_entries))

            # 2. Squad block (all squad duties listed together)
            if squad_by_date:
                squad_entries = []
                for dt in sorted(squad_by_date.keys()):
                    sorted_slots = sorted(squad_by_date[dt], key=lambda x: x[0])
                    time_str = '<br/>'.join([format_time_slot(st, et) for st, et in sorted_slots])
                    squad_entries.append((dt.strftime('%d.%m.%Y'), time_str))
                duty_blocks.append(('Squad', squad_entries))

            if not duty_blocks:
                continue

            for duty_name, date_entries in duty_blocks:
                b_start = cur_f_row
                for d_date, d_time in date_entries:
                    is_fac_first = (cur_f_row == f_start)
                    is_block_first = (cur_f_row == b_start)

                    f_table_rows.append([
                        sl_p if is_fac_first else '',
                        name_p if is_fac_first else '',
                        Paragraph(d_date, cell_norm),
                        Paragraph(d_time, cell_norm),
                        Paragraph(duty_name, cell_norm) if is_block_first else '',
                        ''
                    ])
                    cur_f_row += 1

                b_end = cur_f_row - 1
                if b_end > b_start:
                    f_table_styles.append(('SPAN', (4, b_start), (4, b_end)))

            f_end = cur_f_row - 1
            if f_end > f_start:
                f_table_styles.append(('SPAN', (0, f_start), (0, f_end)))
                f_table_styles.append(('SPAN', (1, f_start), (1, f_end)))
                f_table_styles.append(('SPAN', (5, f_start), (5, f_end)))

        t_fac = Table(f_table_rows, colWidths=[40, 140, 68, 128, 73, 73])
        t_fac.setStyle(TableStyle(f_table_styles))
        elements.append(t_fac)

        # Dynamic spacer calculation so signatories sit gracefully near the bottom on Page 3
        _, h_head3 = header_p3.wrap(522, 800)
        _, t_fac_h = t_fac.wrap(522, 800)
        sig_p3 = ReportService._build_three_signatories_table()
        _, h_sig3 = sig_p3.wrap(522, 800)
        p3_used = h_head3 + 4 + t_fac_h
        p3_push = max(20, min(360, 740 - p3_used - h_sig3))
        elements.append(Spacer(1, p3_push))
        elements.append(sig_p3)

        doc.build(elements)
        output.seek(0)
        return output

    @staticmethod
    def generate_cie_excel(cie, duties, selected_semester=None):
        """
        Generates an official multi-sheet Excel workbook for CIE Duty Allocation matching the 3 pages:
        Sheet 1: Invigilation Allotment
        Sheet 2: Squad Duty Allotment
        Sheet 3: Faculty Roster (with Signatures)
        """
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Invigilation Allotment"

        title_font = Font(name="Calibri", size=13, bold=True)
        sub_font = Font(name="Calibri", size=10, bold=True)
        header_font = Font(name="Calibri", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )

        def write_headers(ws, title, subtitle):
            ws.merge_cells("A1:E1")
            ws["A1"] = "MALNAD COLLEGE OF ENGINEERING, HASSAN"
            ws["A1"].font = title_font
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("A2:E2")
            ws["A2"] = "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING (AI & ML)"
            ws["A2"].font = sub_font
            ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("A3:E3")
            ws["A3"] = f"{cie.name} — {title} ({cie.academic_year})"
            ws["A3"].font = sub_font
            ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

        # Sheet 1: Invigilation
        write_headers(ws1, "Invigilation Duty Allotment", "")
        ws1.append([])
        ws1.append(["Date & Day", "Time Slot", "Course Code", "Room No.", "Faculty (Initials)"])
        r_num = 5
        for col in range(1, 6):
            cell = ws1.cell(row=r_num, column=col)
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        from models import ExamSession, Semester, Faculty
        active_facs = Faculty.query.filter_by(is_active=True).all()
        active_dir_ids = {f.id for f in active_facs}

        s_query = ExamSession.query.filter_by(cie_id=cie.id)
        if selected_semester:
            s_query = s_query.filter_by(semester_id=selected_semester.id)
        else:
            s_query = s_query.join(Semester).filter(Semester.sem_number != 3)
        sessions = s_query.order_by(ExamSession.exam_date, ExamSession.start_time).all()

        for s in sessions:
            r_num += 1
            invig_duties = [d for d in s.duties if (d.duty_type or '').lower() != 'squad duty' and d.faculty_id in active_dir_ids and d.faculty and d.faculty.is_active]
            fac_inits = [get_faculty_initials(d.faculty.name) for d in invig_duties if d.faculty]
            fac_str = ', '.join(fac_inits) if fac_inits else 'TBD'
            rooms_str = ', '.join([r.room_number for r in s.exam_rooms]) or s.room_number or 'AI301'

            row_vals = [
                f"{s.exam_date.strftime('%d.%m.%Y')} ({s.exam_date.strftime('%a')})",
                format_time_slot(s.start_time, s.end_time),
                s.subject.code if s.subject else '',
                rooms_str,
                fac_str
            ]
            for col_idx, val in enumerate(row_vals, 1):
                c = ws1.cell(row=r_num, column=col_idx, value=val)
                c.border = thin_border
                c.alignment = Alignment(horizontal="center", vertical="center")

        # Sheet 2: Squad
        ws2 = wb.create_sheet(title="Squad Duty Allotment")
        write_headers(ws2, "Squad Duty Allotment", "")
        ws2.append([])
        ws2.append(["Date & Day", "Time Slot", "Course Code", "Room No.", "Squad Faculty"])
        r_num2 = 5
        for col in range(1, 6):
            c = ws2.cell(row=r_num2, column=col)
            c.font = header_font
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", vertical="center")

        for s in sessions:
            r_num2 += 1
            squad_duties = [d for d in s.duties if (d.duty_type or '').lower() == 'squad duty' and d.faculty_id in active_dir_ids and d.faculty and d.faculty.is_active]
            s_inits = [get_faculty_initials(d.faculty.name) for d in squad_duties if d.faculty]
            s_str = ', '.join(s_inits) if s_inits else ''
            rooms_str = ', '.join([r.room_number for r in s.exam_rooms]) or s.room_number or 'AI301'

            row_vals = [
                f"{s.exam_date.strftime('%d.%m.%Y')} ({s.exam_date.strftime('%a')})",
                format_time_slot(s.start_time, s.end_time),
                s.subject.code if s.subject else '',
                rooms_str,
                s_str
            ]
            for col_idx, val in enumerate(row_vals, 1):
                c = ws2.cell(row=r_num2, column=col_idx, value=val)
                c.border = thin_border
                c.alignment = Alignment(horizontal="center", vertical="center")

        # Sheet 3: Faculty Roster
        ws3 = wb.create_sheet(title="Faculty Roster")
        ws3.merge_cells("A1:F1")
        ws3["A1"] = f"{cie.name} DUTY ALLOTMENT & SIGN-OFF ({cie.academic_year})"
        ws3["A1"].font = title_font
        ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws3.append([])
        ws3.append(["Sl No.", "Name of the Faculty", "Date", "Time", "Duty", "Signature"])
        r_num3 = 3
        for col in range(1, 7):
            c = ws3.cell(row=r_num3, column=col)
            c.font = header_font
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", vertical="center")

        from models import Duty, Semester
        d_query = Duty.query.join(ExamSession).join(Semester).filter(ExamSession.cie_id == cie.id)
        if selected_semester:
            d_query = d_query.filter(ExamSession.semester_id == selected_semester.id)
        else:
            d_query = d_query.filter(Semester.sem_number != 3)
        fac_duties = d_query.order_by(ExamSession.exam_date, ExamSession.start_time).all()

        faculty_duties_map = defaultdict(list)
        for d in fac_duties:
            if d.faculty_id in active_dir_ids and d.faculty and d.faculty.is_active:
                faculty_duties_map[d.faculty.name].append(d)

        def fac_sort_key_xl(name):
            if 'Arjun' in name: return 0
            if 'Swathi' in name: return 1
            if 'Sushma' in name: return 2
            if 'Ankitha' in name: return 3
            if 'Megha' in name: return 4
            if 'Maseeha' in name: return 5
            if 'Sheethal' in name: return 6
            return 9

        sorted_fac_entries = sorted(faculty_duties_map.items(), key=lambda x: fac_sort_key_xl(x[0]))

        cur_xl_row = 4
        for idx, (fac_name, f_duties) in enumerate(sorted_fac_entries, 1):
            f_start = cur_xl_row
            invig_duties = [d for d in f_duties if (d.duty_type or '').lower() != 'squad duty']
            squad_duties = [d for d in f_duties if (d.duty_type or '').lower() == 'squad duty']

            invig_by_date = defaultdict(list)
            for d in invig_duties:
                sess = d.exam_session
                invig_by_date[sess.exam_date].append((sess.start_time, sess.end_time))

            squad_by_date = defaultdict(list)
            for d in squad_duties:
                sess = d.exam_session
                squad_by_date[sess.exam_date].append((sess.start_time, sess.end_time))

            duty_blocks = []
            if invig_by_date:
                invig_entries = []
                for dt in sorted(invig_by_date.keys()):
                    sorted_slots = sorted(invig_by_date[dt], key=lambda x: x[0])
                    time_str = '\n'.join([format_time_slot(st, et) for st, et in sorted_slots])
                    invig_entries.append((dt.strftime('%d.%m.%Y'), time_str))
                duty_blocks.append(('Invigilation', invig_entries))

            if squad_by_date:
                squad_entries = []
                for dt in sorted(squad_by_date.keys()):
                    sorted_slots = sorted(squad_by_date[dt], key=lambda x: x[0])
                    time_str = '\n'.join([format_time_slot(st, et) for st, et in sorted_slots])
                    squad_entries.append((dt.strftime('%d.%m.%Y'), time_str))
                duty_blocks.append(('Squad', squad_entries))

            if not duty_blocks:
                continue

            for duty_name, date_entries in duty_blocks:
                b_start = cur_xl_row
                for d_date, d_time in date_entries:
                    is_fac_first = (cur_xl_row == f_start)
                    is_block_first = (cur_xl_row == b_start)

                    row_vals = [
                        idx if is_fac_first else "",
                        fac_name if is_fac_first else "",
                        d_date,
                        d_time,
                        duty_name if is_block_first else "",
                        ""
                    ]
                    for col_idx, val in enumerate(row_vals, 1):
                        c = ws3.cell(row=cur_xl_row, column=col_idx, value=val)
                        c.border = thin_border
                        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cur_xl_row += 1

                b_end = cur_xl_row - 1
                if b_end > b_start:
                    ws3.merge_cells(start_row=b_start, start_column=5, end_row=b_end, end_column=5)

            f_end = cur_xl_row - 1
            if f_end > f_start:
                ws3.merge_cells(start_row=f_start, start_column=1, end_row=f_end, end_column=1)
                ws3.merge_cells(start_row=f_start, start_column=2, end_row=f_end, end_column=2)
                ws3.merge_cells(start_row=f_start, start_column=6, end_row=f_end, end_column=6)

        for ws in [ws1, ws2, ws3]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_faculty_duty_slips_pdf(cie, duties, selected_semester=None):
        """
        Generates individual printable duty orders for faculty members.
        Organized portrait per faculty member with official 3 signatories.
        """
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=portrait(letter),
            leftMargin=40,
            rightMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        elements = []
        styles = getSampleStyleSheet()

        inst_title = ParagraphStyle('InstP', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=13, leading=16, alignment=1)
        dept_title = ParagraphStyle('DeptP', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=10, leading=13, alignment=1)
        slip_title = ParagraphStyle('SlipP', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=14, alignment=1)
        body_style = ParagraphStyle('BodyP', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13)
        cell_style = ParagraphStyle('CellP', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=10, alignment=1)
        cell_bold = ParagraphStyle('CellBP', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, leading=10, alignment=1)

        from models import Faculty
        active_facs = Faculty.query.filter_by(is_active=True).all()
        active_dir_ids = {f.id for f in active_facs}

        # Group duties by active directory faculty
        faculty_map = defaultdict(list)
        for d in duties:
            if d.faculty_id in active_dir_ids and d.faculty and d.faculty.is_active:
                faculty_map[d.faculty].append(d)

        for idx, (fac, fac_duties) in enumerate(faculty_map.items()):
            if idx > 0:
                elements.append(PageBreak())

            elements.append(Paragraph("Malnad College of Engineering, Hassan", inst_title))
            elements.append(Paragraph("Department of Computer Science & Engineering (AI & ML)", dept_title))
            elements.append(Spacer(1, 4))
            
            sem_slip_str = f" • {selected_semester.name.upper()}" if selected_semester else f" ({cie.academic_year})"
            elements.append(Paragraph(f"EXAMINATION DUTY ORDER — {cie.name.upper()}{sem_slip_str}", slip_title))
            elements.append(Spacer(1, 12))

            salutation = (
                f"<b>To:</b><br/>"
                f"<b>{fac.name}</b><br/>"
                f"{fac.designation or 'Faculty'}, Department of CSE (AI & ML)<br/>"
                f"Malnad College of Engineering, Hassan"
            )
            elements.append(Paragraph(salutation, body_style))
            elements.append(Spacer(1, 10))

            memo_text = (
                f"You are hereby appointed as an <b>Invigilator</b> for the Continuous Internal Evaluation "
                f"({cie.name}) examinations scheduled between {cie.start_date.strftime('%d %b %Y')} and "
                f"{cie.end_date.strftime('%d %b %Y')}. You are required to strictly report 15 minutes prior to commencement:"
            )
            elements.append(Paragraph(memo_text, body_style))
            elements.append(Spacer(1, 12))

            t_data = [[
                Paragraph("<b>#</b>", cell_bold),
                Paragraph("<b>Date & Day</b>", cell_bold),
                Paragraph("<b>Time Slot</b>", cell_bold),
                Paragraph("<b>Semester</b>", cell_bold),
                Paragraph("<b>Subject Code</b>", cell_bold),
                Paragraph("<b>Room</b>", cell_bold),
                Paragraph("<b>Duty</b>", cell_bold),
            ]]

            for s_idx, fd in enumerate(fac_duties, 1):
                sess = fd.exam_session
                room_str = fd.exam_room.room_number if fd.exam_room else (sess.room_number or 'AI301')
                d_type = 'Squad' if (fd.duty_type or '').lower() == 'squad duty' else 'Invigilation'
                t_data.append([
                    Paragraph(str(s_idx), cell_style),
                    Paragraph(f"{sess.exam_date.strftime('%d.%m.%Y')}<br/>{sess.exam_date.strftime('%a')}", cell_style),
                    Paragraph(format_time_slot(sess.start_time, sess.end_time), cell_style),
                    Paragraph(sess.semester.name if sess.semester else "", cell_style),
                    Paragraph(sess.subject.code if sess.subject else "", cell_style),
                    Paragraph(f"<b>{room_str}</b>", cell_style),
                    Paragraph(d_type, cell_style),
                ])

            duty_table = Table(t_data, colWidths=[24, 90, 130, 80, 80, 56, 70])
            duty_table.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            elements.append(duty_table)
            elements.append(Spacer(1, 20))

            elements.append(ReportService._build_three_signatories_table())

        doc.build(elements)
        output.seek(0)
        return output
